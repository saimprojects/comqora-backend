"""Bounded, evidence-first extraction for courier payment statements.

The importer never creates a ledger entry from a file. It turns a PDF, CSV or
spreadsheet into a reviewable draft with raw source values, table headings and
parser evidence. Financial posting remains behind the reconciliation checks in
``settlements.py``.
"""

import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

import pdfplumber

from .courier_profiles import normalize_label, profile_for, role_override

MAX_BYTES = 8 * 1024 * 1024
MAX_PAGES = 25
MAX_ROWS = 1000
MAX_SHEETS = 12
MAX_COLUMNS = 30
MAX_SPREADSHEET_CELLS = 50000
ROLES = {
    "tracking",
    "order_ref",
    "gross",
    "net",
    "fee",
    "deduction",
    "credit",
    "fee_credit",
    "info",
    "unknown",
}
IDENTIFIER_ROLES = {"tracking", "order_ref"}
FINANCIAL_ROLES = {"gross", "net", "fee", "deduction", "credit", "fee_credit"}
PDF_MAGIC = b"%PDF-"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def layout_signature(columns):
    labels = [" ".join(col["label"].lower().split()) for col in columns]
    return hashlib.sha256(json.dumps(labels).encode()).hexdigest()


def amount(value):
    """Strict PKR money parsing; missing and ambiguous values are never zero."""
    if isinstance(value, bool) or not isinstance(value, (str, int, Decimal)):
        raise ValueError("Enter an amount with at most two decimal places.")
    text = str(value).strip().replace("−", "-").replace("\u00a0", " ")
    text = re.sub(r"^(?:PKR|Rs\.?|Rupees?)\s*", "", text, flags=re.I)
    text = re.sub(r"\s*/-\s*$", "", text)
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1].strip()
    if not re.fullmatch(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?", text):
        raise ValueError("Use 1234.56 format; blanks are not zero.")
    try:
        result = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError("Invalid amount.") from exc
    if abs(result) > Decimal("9999999999.99"):
        raise ValueError("Amount is too large.")
    return result.quantize(Decimal(".01"))


def infer_role(label, profile=None):
    """Map common courier headings without treating vague fields as money facts."""
    key = normalize_label(label)
    override = role_override(label, profile or {})
    if override:
        return override
    # Weight and payment can both contain "net". Classify the former first.
    if re.search(r"\b(net|actual|charged)? ?weight\b|\bweight kg\b", key):
        return "info"
    if re.search(r"\b(tracking|consignment|waybill|airway bill|awb)\b|^cn(?: no| number)?$", key):
        return "tracking"
    if re.search(
        r"\b(order|merchant) (?:ref(?:erence)?|id|no|number)\b|\binvoice (?:no|number|id)\b",
        key,
    ):
        return "order_ref"
    if re.search(
        r"\b(net amount|net payable|payable amount|amount payable|remittance amount|net settlement|net payment)\b|^(?:net|payable)$",
        key,
    ):
        return "net"
    if re.search(
        r"\b(?:cod(?: (?:amount|value))?|collected(?: amount)?|collection(?: amount)?)\b|\bgross amount\b",
        key,
    ):
        return "gross"
    if re.search(
        r"\b(?:delivery|shipping|freight|postage)(?: (?:charge|charges|fee|fees))?\b|\b(?:cash handling|service|cod) (?:charge|charges|fee|fees)\b",
        key,
    ):
        return "fee"
    if re.search(r"\b(deduction|tax deduction)\b", key):
        return "deduction"
    # A bare GST/VAT can be a courier cost, an income withholding or a total.
    # Force an explicit review unless a courier profile gives it a known meaning.
    if re.search(
        r"\b(tax|gst|vat|wht|withholding|adjustment|reserve|carry forward|upfront)\b", key
    ):
        return "unknown"
    if re.search(
        r"\b(sr|s no|sno|serial|city|status|date|weight|name|address|phone|reference|ref|origin|destination|pickup|delivery)\b",
        key,
    ):
        return "info"
    return "unknown"


def _first_line(raw):
    return str(raw or "").strip().replace("\u00a0", " ").split("\n")[0].strip()


def tracking_value(raw):
    """Keep a carrier identifier, avoiding dates/amounts in a multi-line CN cell."""
    first = _first_line(raw)
    return (
        first
        if re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9/_-]{3,119}", first)
        and any(char.isdigit() for char in first)
        else ""
    )


def order_reference_value(raw):
    first = _first_line(raw)
    if not first or len(first) > 120 or not re.search(r"[A-Za-z0-9]", first):
        return ""
    return first


def cell_value(raw, role):
    raw = str(raw or "").strip()
    if role == "tracking":
        return tracking_value(raw)
    if role == "order_ref":
        return order_reference_value(raw)
    if role in FINANCIAL_ROLES:
        try:
            # Parentheses in a charge column represent an outflow, not a second
            # negative deduction. Credits need an explicit credit role.
            return str(amount(raw) if role in {"gross", "net"} else abs(amount(raw)))
        except ValueError:
            return raw
    return raw


def _is_identifier(raw, role):
    return bool(tracking_value(raw) if role == "tracking" else order_reference_value(raw))


def _header_tables(page, profile):
    """Find ruled headers first, then borderless tables using text alignment."""
    for strategy in [
        None,
        {"vertical_strategy": "text", "horizontal_strategy": "text", "min_words_vertical": 2},
    ]:
        candidates = []
        for table in page.find_tables(strategy or {}):
            extracted = table.extract()
            for index, cells in enumerate(extracted[:10]):
                labels = [str(cell or "").strip() for cell in cells]
                roles = [infer_role(cell, profile) for cell in labels]
                if not any(role in IDENTIFIER_ROLES for role in roles) or not any(
                    role in {"gross", "net", "fee"} for role in roles
                ):
                    continue
                boxes = table.rows[index].cells
                # Text-alignment detection sometimes inserts blank gutters. Drop
                # only truly empty ones; an unlabeled data column must remain.
                keep = [
                    column
                    for column, label in enumerate(labels)
                    if label
                    or any(
                        str(row[column] or "").strip()
                        for row in extracted[index + 1 :]
                        if column < len(row)
                    )
                ]
                labels = [labels[column] or f"Unlabelled {column + 1}" for column in keep]
                boxes = [boxes[column] for column in keep]
                if all(boxes) and len(labels) <= MAX_COLUMNS:
                    candidates.append((labels, boxes, table.bbox))
                break
        if candidates:
            return candidates
    return []


def _pdf_rows(words, labels, boxes, table_box, top, next_top, page_number, profile):
    roles = [infer_role(label, profile) for label in labels]
    identifier_index = next(index for index, role in enumerate(roles) if role in IDENTIFIER_ROLES)
    identifier_role = roles[identifier_index]
    identifier_box = boxes[identifier_index]
    anchors = [
        word
        for word in words
        if top <= word["top"] < next_top
        and identifier_box[0] <= (word["x0"] + word["x1"]) / 2 <= identifier_box[2]
        and _is_identifier(word["text"], identifier_role)
        # An identifier can contain a slash; only obvious decimal amounts are
        # excluded. This fixes CNs such as ABC/00124.
        and not re.fullmatch(r"\d[\d,]*\.\d{1,2}", word["text"])
    ]
    rows = []
    for index, anchor in enumerate(anchors):
        row_top = anchor["top"] - 3
        row_bottom = (
            anchors[index + 1]["top"] - 3
            if index + 1 < len(anchors)
            else min(next_top, anchor["bottom"] + 38)
        )
        band = [word for word in words if row_top <= word["top"] < row_bottom]
        money_words = [
            word
            for word in band
            if re.fullmatch(r"\(?-?\d[\d,]*\.\d{2}\)?", word["text"])
            and word["x0"] >= boxes[min(identifier_index + 1, len(boxes) - 1)][0]
        ]
        baseline = (
            min(money_words, key=lambda word: abs(word["top"] - anchor["top"]))["top"]
            if money_words
            else anchor["top"]
        )
        raw = []
        for column, box in enumerate(boxes):
            selected = [word for word in band if box[0] <= (word["x0"] + word["x1"]) / 2 < box[2]]
            if (
                roles[column] in FINANCIAL_ROLES
                or roles[column] == "unknown"
                or "amount" in labels[column].lower()
            ):
                selected = [word for word in selected if abs(word["top"] - baseline) < 5]
            if column == identifier_index:
                raw.append(anchor["text"])
            else:
                raw.append(
                    " ".join(
                        word["text"]
                        for word in sorted(
                            selected, key=lambda word: (round(word["top"] / 3), word["x0"])
                        )
                    )
                )
        rows.append(
            {
                "values": [cell_value(value, role) for value, role in zip(raw, roles)],
                "raw": raw,
                "page": page_number,
                "bbox": [
                    round(table_box[0], 2),
                    round(row_top, 2),
                    round(table_box[2], 2),
                    round(row_bottom, 2),
                ],
                "external": False,
            }
        )
    return rows


def _result(document_kind, profile_id, profile_evidence):
    return {
        "version": 2,
        "document_kind": document_kind,
        "pages": [],
        "tables": [],
        "summary_candidates": [],
        "warnings": [],
        "parser": {
            "profile": profile_id,
            "profile_evidence": profile_evidence,
            "confidence": 0,
        },
    }


def _finish_result(result):
    tables = result["tables"]
    rows = sum(len(table["rows"]) for table in tables)
    known_columns = sum(
        1
        for table in tables
        for column in table["columns"]
        if column["role"] not in {"unknown", "info"}
    )
    unknown_columns = sum(
        1 for table in tables for column in table["columns"] if column["role"] == "unknown"
    )
    confidence = 0
    if tables:
        confidence = min(99, 42 + min(rows, 30) + min(known_columns * 3, 25) - unknown_columns * 2)
    result["parser"].update(
        {
            "confidence": max(0, confidence),
            "row_count": rows,
            "unknown_columns": unknown_columns,
        }
    )
    if rows > MAX_ROWS:
        raise ValueError("Use a statement with at most 1,000 shipment rows.")
    if not tables:
        result["warnings"].append(
            "No reliable shipment table was detected. Add columns and rows in the review; the original source remains available."
        )
    elif unknown_columns:
        result["warnings"].append(
            f"{unknown_columns} column(s) need a human classification before confirmation."
        )
    result["warnings"].append(
        "Column roles and totals are suggestions. Verify all pages, rows, deductions and payment totals before confirming."
    )
    return result


def extract_pdf(blob, courier_hint=""):
    if not blob.startswith(PDF_MAGIC) or len(blob) > MAX_BYTES:
        raise ValueError("Upload a PDF of at most 8 MB.")
    profile_id, profile, profile_evidence = profile_for(courier_hint)
    result = _result("pdf", profile_id, profile_evidence)
    previous_headers = []
    with pdfplumber.open(io.BytesIO(blob)) as pdf:
        if not 1 <= len(pdf.pages) <= MAX_PAGES:
            raise ValueError("Use a PDF containing 1–25 pages.")
        for page_number, page in enumerate(pdf.pages, 1):
            if len(page.chars) > 150000 or page.width > 3000 or page.height > 4000:
                raise ValueError("This page exceeds the safe processing limit.")
            words = page.extract_words(x_tolerance=2, y_tolerance=3)
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            if len(text) > 80000:
                raise ValueError("This page contains too much text.")
            result["pages"].append(
                {"page": page_number, "text": text, "width": page.width, "height": page.height}
            )
            if not words:
                result["warnings"].append(
                    f"Page {page_number}: no readable text. This may be a scan; use a text-based export or transcribe it against the preview."
                )
                previous_headers = []
                page.close()
                continue
            detected_headers = _header_tables(page, profile)
            headers = detected_headers
            if not headers and previous_headers:
                # Portal statements commonly print a header only on page one. A
                # continuation is evidence-marked and still requires review.
                headers = previous_headers
                result["warnings"].append(
                    f"Page {page_number}: reused the previous page's table headings for a continuation."
                )
            for header_index, (labels, boxes, table_box) in enumerate(headers):
                top = max(box[3] for box in boxes)
                next_top = (
                    headers[header_index + 1][1][0][1]
                    if header_index + 1 < len(headers)
                    else page.height
                )
                rows = _pdf_rows(
                    words,
                    labels,
                    boxes,
                    table_box,
                    top,
                    next_top,
                    page_number,
                    profile,
                )
                if rows:
                    result["tables"].append(
                        {
                            "columns": [
                                {"label": label, "role": infer_role(label, profile)}
                                for label in labels
                            ],
                            "rows": rows,
                            "page": page_number,
                        }
                    )
            if detected_headers:
                previous_headers = detected_headers
            for line in text.splitlines():
                if re.search(
                    r"total|payable|settlement|deduction|gst|tax|charge|balance|adjustment|carry|cpr",
                    line,
                    re.I,
                ):
                    result["summary_candidates"].append({"page": page_number, "text": line})
            page.close()
    # If the selected courier uses an arbitrary name but its document advertises
    # a known profile, re-classify headings after reading the evidence.
    if profile_id == "generic":
        source_profile, source_data, source_evidence = profile_for(
            "", "\n".join(page["text"] for page in result["pages"])
        )
        if source_profile != "generic":
            result["parser"].update(profile=source_profile, profile_evidence=source_evidence)
            for table in result["tables"]:
                for column in table["columns"]:
                    column["role"] = infer_role(column["label"], source_data)
    return _finish_result(result)


def _spreadsheet_cell(value):
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).strip()


def _tabular_header(rows, profile):
    for row_index, row in enumerate(rows[:40]):
        labels = [value.strip() or f"Unlabelled {index + 1}" for index, value in enumerate(row)]
        if not labels or len(labels) > MAX_COLUMNS:
            continue
        roles = [infer_role(label, profile) for label in labels]
        if any(role in IDENTIFIER_ROLES for role in roles) and any(
            role in {"gross", "net", "fee"} for role in roles
        ):
            return row_index, labels, roles
    return None


def _tabular_result(sheets, document_kind, courier_hint=""):
    joined = "\n".join("\n".join(" | ".join(row) for row in rows[:80]) for _, rows in sheets)
    profile_id, profile, profile_evidence = profile_for(courier_hint, joined)
    result = _result(document_kind, profile_id, profile_evidence)
    total_cells = 0
    for page_number, (sheet_name, raw_rows) in enumerate(sheets, 1):
        if page_number > MAX_SHEETS:
            result["warnings"].append(f"Only the first {MAX_SHEETS} sheets were examined.")
            break
        rows = [row[:MAX_COLUMNS] for row in raw_rows]
        total_cells += sum(len(row) for row in rows)
        if total_cells > MAX_SPREADSHEET_CELLS:
            raise ValueError("The spreadsheet contains too many cells for a safe statement review.")
        text = "\n".join(" | ".join(row) for row in rows[:500])[:80000]
        result["pages"].append({"page": page_number, "sheet": sheet_name[:100], "text": text})
        detected = _tabular_header(rows, profile)
        if not detected:
            result["warnings"].append(
                f"Sheet '{sheet_name[:80]}' has no reliable shipment header and was not auto-mapped."
            )
            continue
        header_index, labels, roles = detected
        identifier_index = next(
            index for index, role in enumerate(roles) if role in IDENTIFIER_ROLES
        )
        table_rows = []
        empty_rows = 0
        for source_row in rows[header_index + 1 :]:
            values = source_row + [""] * (len(labels) - len(source_row))
            values = values[: len(labels)]
            if not any(value.strip() for value in values):
                empty_rows += 1
                if empty_rows >= 3 and table_rows:
                    break
                continue
            empty_rows = 0
            identifier = values[identifier_index].strip()
            if not _is_identifier(identifier, roles[identifier_index]):
                # Totals, notes and section headings cannot become shipments.
                continue
            if not any(
                values[index].strip() for index, role in enumerate(roles) if role in FINANCIAL_ROLES
            ):
                continue
            table_rows.append(
                {
                    "values": [cell_value(value, role) for value, role in zip(values, roles)],
                    "raw": values,
                    "page": page_number,
                    "sheet": sheet_name[:100],
                    "source_row": header_index + len(table_rows) + 2,
                    "external": False,
                }
            )
            if sum(len(table["rows"]) for table in result["tables"]) + len(table_rows) > MAX_ROWS:
                raise ValueError("Use a statement with at most 1,000 shipment rows.")
        if table_rows:
            result["tables"].append(
                {
                    "columns": [
                        {"label": label, "role": role} for label, role in zip(labels, roles)
                    ],
                    "rows": table_rows,
                    "page": page_number,
                    "sheet": sheet_name[:100],
                }
            )
        for row in rows[: min(len(rows), 150)]:
            line = " | ".join(row).strip()
            if re.search(
                r"total|payable|settlement|deduction|gst|tax|charge|balance|adjustment|carry|cpr",
                line,
                re.I,
            ):
                result["summary_candidates"].append({"page": page_number, "text": line})
    return _finish_result(result)


def _safe_xlsx(blob):
    try:
        with ZipFile(io.BytesIO(blob)) as archive:
            infos = archive.infolist()
            if len(infos) > 500 or sum(item.file_size for item in infos) > 48 * 1024 * 1024:
                raise ValueError("The spreadsheet expands beyond the safe processing limit.")
            if any(
                item.compress_size and item.file_size / item.compress_size > 150 for item in infos
            ):
                raise ValueError("The spreadsheet compression ratio is unsafe.")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("Upload a valid .xlsx spreadsheet.")
    except BadZipFile as exc:
        raise ValueError("Upload a valid .xlsx spreadsheet.") from exc


def extract_xlsx(blob, courier_hint=""):
    _safe_xlsx(blob)
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - deployment dependency guard
        raise ValueError("Spreadsheet support is not installed on this server.") from exc
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(blob), read_only=True, data_only=True, keep_links=False
        )
        sheets = []
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            rows = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                if row_index > MAX_ROWS + 80:
                    break
                rows.append([_spreadsheet_cell(value) for value in row[:MAX_COLUMNS]])
            sheets.append((worksheet.title, rows))
        workbook.close()
    except Exception as exc:
        raise ValueError("The .xlsx statement could not be read safely.") from exc
    return _tabular_result(sheets, "xlsx", courier_hint)


def extract_xls(blob, courier_hint=""):
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - deployment dependency guard
        raise ValueError("Legacy .xls support is not installed on this server.") from exc
    try:
        workbook = xlrd.open_workbook(file_contents=blob, on_demand=True)
        sheets = []
        for index in range(min(workbook.nsheets, MAX_SHEETS)):
            sheet = workbook.sheet_by_index(index)
            rows = []
            for row_index in range(min(sheet.nrows, MAX_ROWS + 80)):
                values = []
                for column in range(min(sheet.ncols, MAX_COLUMNS)):
                    cell = sheet.cell(row_index, column)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    values.append(_spreadsheet_cell(value))
                rows.append(values)
            sheets.append((sheet.name, rows))
        workbook.release_resources()
    except Exception as exc:
        raise ValueError("The .xls statement could not be read safely.") from exc
    return _tabular_result(sheets, "xls", courier_hint)


def extract_delimited(blob, filename="", courier_hint=""):
    decoded = ""
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            decoded = blob.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not decoded or "\x00" in decoded:
        raise ValueError("Upload a UTF-8, UTF-16 or Windows-encoded CSV/TSV statement.")
    if len(decoded) > MAX_BYTES:
        raise ValueError("Statement is too large.")
    try:
        dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if PurePath(filename).suffix.lower() == ".tsv" else csv.excel
    reader = csv.reader(io.StringIO(decoded, newline=""), dialect)
    rows = []
    try:
        for row_index, row in enumerate(reader, 1):
            if row_index > MAX_ROWS + 80:
                break
            rows.append([str(value).strip() for value in row[:MAX_COLUMNS]])
    except csv.Error as exc:
        raise ValueError("The CSV/TSV statement is malformed.") from exc
    return _tabular_result(
        [(PurePath(filename).stem or "Statement", rows)], "delimited", courier_hint
    )


def detect_document_kind(blob, filename=""):
    if not blob or len(blob) > MAX_BYTES:
        raise ValueError("Upload a statement of at most 8 MB.")
    if blob.startswith(PDF_MAGIC):
        return "pdf"
    if blob.startswith(OLE_MAGIC):
        return "xls"
    if blob.startswith(b"PK\x03\x04"):
        _safe_xlsx(blob)
        return "xlsx"
    suffix = PurePath(filename).suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        return "delimited"
    raise ValueError("Upload an original PDF, .xlsx, .xls, .csv or .tsv courier statement.")


def source_content_type(blob, filename=""):
    return {
        "pdf": "application/pdf",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xls": "application/vnd.ms-excel",
        "delimited": "text/csv; charset=utf-8",
    }[detect_document_kind(blob, filename)]


def extract_document(blob, filename="", courier_hint=""):
    kind = detect_document_kind(blob, filename)
    if kind == "pdf":
        return extract_pdf(blob, courier_hint)
    if kind == "xlsx":
        return extract_xlsx(blob, courier_hint)
    if kind == "xls":
        return extract_xls(blob, courier_hint)
    return extract_delimited(blob, filename, courier_hint)


_DATE_TOKEN = (
    r"(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}|"
    r"\d{1,2}\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4})"
)


def _parse_date(raw):
    raw = re.sub(r"\s+", " ", raw.strip())
    formats = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d-%m-%Y", "%d.%m.%Y")
    if re.search(r"[A-Za-z]", raw):
        formats = ("%d %b %Y", "%d %B %Y")
    elif "/" in raw:
        # A 10/09 date is genuinely ambiguous. Do not guess a US/PK order.
        fields = raw.split("/")
        if len(fields) == 3 and int(fields[0]) <= 12 and int(fields[1]) <= 12:
            return "", True
        formats = ("%d/%m/%Y", "%d/%m/%y")
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).date().isoformat(), False
        except ValueError:
            continue
    return "", False


def _metadata(text, profile):
    refs, dates, nets, gross = [], [], [], []
    reference_terms = [
        "cpr",
        "settlement",
        "remittance",
        "payment",
        "voucher",
        "disbursement",
        *profile.get("reference_terms", ()),
    ]
    ref_terms = "|".join(
        re.escape(term) for term in sorted(set(reference_terms), key=len, reverse=True)
    )
    ref_pattern = re.compile(
        rf"(?:{ref_terms})\s*(?:number|no\.?|ref(?:erence)?|id|#)?\s*[:#-]?\s*"
        r"([A-Za-z0-9][A-Za-z0-9/_-]{3,119})",
        re.I,
    )
    for match in ref_pattern.finditer(text):
        candidate = match.group(1)
        if not re.fullmatch(r"(?:date|amount|total|status)", candidate, re.I):
            refs.append(candidate)
    for match in re.finditer(
        rf"(?:statement|settlement|remittance|payment|cpr)\s*date\s*[:#-]?\s*{_DATE_TOKEN}",
        text,
        re.I,
    ):
        parsed, ambiguous = _parse_date(match.group(1))
        if parsed:
            dates.append(parsed)
        elif ambiguous:
            dates.append("AMBIGUOUS:" + match.group(1))
    lines = text.splitlines()
    for index, line in enumerate(lines):
        normal = normalize_label(line)
        values = re.findall(r"\(?-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?\)?", line)
        if not values and index + 1 < len(lines):
            values = re.findall(
                r"^\s*(\(?-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?\)?)\s*$",
                lines[index + 1],
            )
        if re.search(
            r"\b(net total|net payable|total payable|net settlement|net payment|final payable)\b",
            normal,
        ):
            for value in values:
                try:
                    nets.append(str(amount(value)))
                except ValueError:
                    pass
        if re.search(r"\b(gross total|total cod|total collection|total collected)\b", normal):
            for value in values:
                try:
                    gross.append(str(amount(value)))
                except ValueError:
                    pass
    return {"references": refs, "dates": dates, "net": nets, "gross": gross}


def _unique(values):
    clean = [value for value in values if value]
    distinct = list(dict.fromkeys(clean))
    return distinct[0] if len(distinct) == 1 else ""


def draft_review(extracted):
    pages = extracted.get("pages", [])
    text = "\n".join(page.get("text", "") for page in pages)
    profile_id = extracted.get("parser", {}).get("profile", "generic")
    _, profile, _ = profile_for(profile_id)
    metadata = _metadata(text, profile)
    extracted["metadata_candidates"] = metadata
    return {
        "reference": _unique(metadata["references"]),
        "date": _unique(metadata["dates"]),
        "currency": "PKR",
        "declared_net": _unique(metadata["net"]),
        "declared_gross": _unique(metadata["gross"]),
        "tables": extracted.get("tables", []),
        "adjustments": [],
        "checks": [],
        "update_costs": False,
        "replace_costs": False,
        "ownership_confirmed": False,
        "source_confirmed": False,
        "notes": "",
    }


def render_page(blob, page_number):
    if detect_document_kind(blob) != "pdf":
        raise ValueError(
            "Page previews are available for PDF statements only. Download the original spreadsheet to inspect it."
        )
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(blob) as pdf:
        if not 1 <= page_number <= min(len(pdf), MAX_PAGES):
            raise ValueError("Page not found.")
        page = pdf[page_number - 1]
        width, height = page.get_size()
        if width <= 0 or height <= 0 or width > 3000 or height > 4000:
            raise ValueError("Page dimensions exceed the preview limit.")
        bitmap = page.render(scale=min(1.6, 1600 / max(width, height)))
        output = io.BytesIO()
        bitmap.to_pil().save(output, format="PNG")
        bitmap.close()
        page.close()
        return output.getvalue()
